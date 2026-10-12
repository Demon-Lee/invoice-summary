import argparse
import json
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


ACCEPTANCE_SHEET = "验收"
ISSUE_SHEET = "领用"
SEQUENCE_HEADER = "序号"
AMOUNT_HEADER = "金额"
TOTAL_MARKER = "合计"
FOOTER_RULES = {
    ACCEPTANCE_SHEET: ("经办人", "验收人"),
    ISSUE_SHEET: ("保管人", "领用人"),
}
CENT = Decimal("0.01")


def _value_text(value):
    return str(value).strip() if value is not None else ""


def _decimal(value, description):
    if isinstance(value, bool) or value is None or not _value_text(value):
        raise ValueError(f"{description} is missing")
    try:
        result = Decimal(_value_text(value).replace(",", ""))
    except InvalidOperation as error:
        raise ValueError(f"{description} is not a numeric value") from error
    cents = result.quantize(CENT)
    if result != cents:
        raise ValueError(f"{description} must be exact to cents")
    return cents


def _sequence_key(value, row_number):
    if value is None or not _value_text(value):
        raise ValueError(f"验收 row {row_number} is missing a 序号 value")
    if isinstance(value, bool):
        raise ValueError(f"验收 row {row_number} has an invalid 序号 value")
    if isinstance(value, (int, float, Decimal)):
        try:
            numeric = Decimal(str(value))
        except InvalidOperation as error:
            raise ValueError(f"验收 row {row_number} has an invalid 序号 value") from error
        return format(numeric.normalize(), "f")
    return _value_text(value)


def _header_columns(sheet):
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        columns = {_value_text(value): index for index, value in enumerate(row) if _value_text(value)}
        if SEQUENCE_HEADER in columns and AMOUNT_HEADER in columns:
            return row_number, columns
    raise ValueError(
        f"{sheet.title} worksheet is missing required {SEQUENCE_HEADER} and {AMOUNT_HEADER} headers"
    )


def _validated_footer_row(sheet, header_row):
    total_rows = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
        if _value_text(row[0] if row else None).startswith(TOTAL_MARKER):
            total_rows.append(row_number)
    if not total_rows:
        raise ValueError(f"工作表“{sheet.title}”缺少合计页脚")
    if len(total_rows) > 1:
        rows = "、".join(str(row_number) for row_number in total_rows)
        raise ValueError(f"工作表“{sheet.title}”检测到多个合计页脚：第 {rows} 行")

    total_row = total_rows[0]
    try:
        primary_signer, secondary_signer = FOOTER_RULES[sheet.title]
    except KeyError as error:
        raise ValueError(f"工作表“{sheet.title}”没有已定义的页脚规则") from error

    signature_row = total_row + 1
    if (
        not _value_text(sheet.cell(signature_row, 1).value).startswith(primary_signer)
        or not _value_text(sheet.cell(signature_row, 4).value).startswith(secondary_signer)
    ):
        raise ValueError(
            f"工作表“{sheet.title}”第 {signature_row} 行签字行不符合要求："
            f"A 列需要“{primary_signer}”，D 列需要“{secondary_signer}”"
        )

    description_row = total_row + 2
    if not _value_text(sheet.cell(description_row, 1).value).startswith("说明"):
        raise ValueError(
            f"工作表“{sheet.title}”第 {description_row} 行说明行不符合要求："
            "A 列需要“说明”"
        )
    return total_row


def _footer_total(sheet, columns, footer_row):
    amount = sheet.cell(footer_row, columns[AMOUNT_HEADER] + 1).value
    description = f"{sheet.title} worksheet final 合计 row 金额 value"
    if isinstance(amount, bool) or amount is None or not _value_text(amount):
        raise ValueError(f"{description} is missing")
    try:
        result = Decimal(_value_text(amount).replace(",", ""))
    except InvalidOperation as error:
        raise ValueError(f"{description} is not a numeric value") from error
    cents = result.quantize(CENT)
    # Formula caches are IEEE-754 numbers.  A cent-exact Excel formula can be
    # serialized as e.g. 1540.1100000000001; accept that insignificant cache
    # tail while still rejecting a real fractional-cent total.
    if abs(result - cents) > Decimal("0.000001"):
        raise ValueError(f"{description} must be exact to cents")
    return cents


def _acceptance_totals(sheet, columns, header_row, footer_row):
    totals = {}
    sequence_column = columns[SEQUENCE_HEADER] + 1
    amount_column = columns[AMOUNT_HEADER] + 1
    for row_number in range(header_row + 1, footer_row):
        sequence = sheet.cell(row_number, sequence_column).value
        amount = sheet.cell(row_number, amount_column).value
        if not _value_text(sequence) and not _value_text(amount):
            continue
        key = _sequence_key(sequence, row_number)
        totals[key] = totals.get(key, Decimal("0.00")) + _decimal(
            amount, f"验收 row {row_number} 金额 value"
        )
    return {sequence: format(total, ".2f") for sequence, total in totals.items()}


def collect_excel_totals(input_path):
    input_path = Path(input_path)
    if input_path.suffix.lower() != ".xlsx":
        raise ValueError("Excel totals collector accepts .xlsx workbooks only")
    try:
        workbook = load_workbook(input_path, read_only=True, data_only=True)
    except Exception as error:
        raise ValueError(f"unable to read workbook: {error}") from error
    try:
        missing = [name for name in (ACCEPTANCE_SHEET, ISSUE_SHEET) if name not in workbook.sheetnames]
        if missing:
            raise ValueError(f"workbook is missing required worksheet: {', '.join(missing)}")
        acceptance = workbook[ACCEPTANCE_SHEET]
        issue = workbook[ISSUE_SHEET]
        acceptance_header, acceptance_columns = _header_columns(acceptance)
        issue_header, issue_columns = _header_columns(issue)
        acceptance_footer = _validated_footer_row(acceptance, acceptance_header)
        issue_footer = _validated_footer_row(issue, issue_header)
        return {
            "sequences": _acceptance_totals(
                acceptance, acceptance_columns, acceptance_header, acceptance_footer
            ),
            "acceptance_footer_total": format(
                _footer_total(acceptance, acceptance_columns, acceptance_footer), ".2f"
            ),
            "issue_footer_total": format(
                _footer_total(issue, issue_columns, issue_footer), ".2f"
            ),
        }
    finally:
        workbook.close()


def _validated_output_path(input_value, output_value):
    if output_value is None:
        return None
    input_path = Path(input_value)
    output_path = Path(output_value)
    if output_path.resolve(strict=False) == input_path.resolve(strict=False):
        raise ValueError("output path must differ from the input workbook")
    if input_path.exists() and output_path.exists() and output_path.samefile(input_path):
        raise ValueError("output path must differ from the input workbook")
    if output_path.suffix.lower() == ".xlsx":
        raise ValueError("output path must not be an .xlsx workbook")
    return output_path


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Collect read-only acceptance and issue totals from an .xlsx workbook."
    )
    parser.add_argument("--input", required=True, help="completed .xlsx workbook to inspect")
    parser.add_argument("--output", help="optional JSON output path")
    arguments = parser.parse_args(argv)
    try:
        output_path = _validated_output_path(arguments.input, arguments.output)
        payload = collect_excel_totals(arguments.input)
    except ValueError as error:
        parser.error(str(error))
    rendered = json.dumps(payload, ensure_ascii=False, indent=2)
    if output_path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(rendered + "\n", encoding="utf-8")
    sys.stdout.write(rendered + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
