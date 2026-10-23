import hashlib
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from invoice_desktop.controller import ExcelAdapter
from tests.fixtures import template


class FormatContentGuardTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='invoice-content-guard-')
        self.addCleanup(self.temporary.cleanup)
        self.path = template(Path(self.temporary.name) / 'sample.xlsx')
        wb = load_workbook(self.path)
        for ws in wb:
            for column, value in enumerate([1, '测试器材', None, '个', 1, 2.286666666666667, 12.34], 1):
                ws.cell(4, column, value)
        wb.save(self.path)
        wb.close()

    def run_format(self, mutate):
        def native_run(script, extra, **options):
            staged = Path(extra[extra.index('-WorkbookInput') + 1])
            mutate(staged)
        with patch.object(ExcelAdapter, '_run', side_effect=native_run):
            ExcelAdapter(backend='WPS').apply(Path('shared-fonts.ps1'), self.path)

    def change_cell(self, coordinate, value):
        def mutate(path):
            wb = load_workbook(path)
            wb['领用'][coordinate] = value
            wb.save(path)
            wb.close()
        return mutate

    def assert_rejected(self, mutate):
        before = hashlib.sha256(self.path.read_bytes()).hexdigest()
        with self.assertRaisesRegex(ValueError, '单元格内容或合并结构发生变化'):
            self.run_format(mutate)
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).hexdigest(), before)

    def assert_accepted(self, mutate):
        try:
            self.run_format(mutate)
        except ValueError as error:
            self.fail(str(error))

    def test_native_empty_text_is_equivalent_to_a_blank_cell(self):
        def add_empty_text(path):
            with zipfile.ZipFile(path) as archive:
                files = {name: archive.read(name) for name in archive.namelist()}
            namespace = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
            tag = lambda name: '{' + namespace + '}' + name
            tree = ET.fromstring(files['xl/worksheets/sheet2.xml'])
            row = next(row for row in tree.find(tag('sheetData')) if row.get('r') == '4')
            cell = ET.SubElement(row, tag('c'), {'r': 'C4', 't': 'inlineStr'})
            ET.SubElement(ET.SubElement(cell, tag('is')), tag('t')).text = ''
            files['xl/worksheets/sheet2.xml'] = ET.tostring(tree, encoding='utf-8')
            with zipfile.ZipFile(path, 'w') as archive:
                for name, data in files.items():
                    archive.writestr(name, data)
            wb = load_workbook(path)
            self.assertEqual(wb['领用']['C4'].value, '')
            wb.close()
        self.assert_accepted(add_empty_text)

    def test_unit_price_roundtrip_at_fifteen_significant_digits_is_accepted(self):
        self.assert_accepted(self.change_cell('F4', 2.28666666666667))

    def test_changed_unit_price_is_rejected(self):
        self.assert_rejected(self.change_cell('F4', 2.28666666666676))

    def test_amount_and_quantity_rounding_are_not_exempt(self):
        for coordinate in ('E4', 'G4'):
            with self.subTest(coordinate=coordinate):
                wb = load_workbook(self.path)
                wb['领用'][coordinate] = 2.286666666666667
                wb.save(self.path)
                wb.close()
                self.assert_rejected(self.change_cell(coordinate, 2.28666666666667))

    def test_numeric_type_changes_and_whitespace_are_rejected(self):
        for coordinate, value in [('F4', '2.286666666666667'), ('C4', ' '), ('B4', '其他器材')]:
            with self.subTest(coordinate=coordinate):
                self.assert_rejected(self.change_cell(coordinate, value))

    def test_changed_formula_is_rejected(self):
        self.assert_rejected(self.change_cell('G7', '=SUM(G4:G5)'))

    def test_changed_merge_is_rejected(self):
        def mutate(path):
            wb = load_workbook(path)
            wb['领用'].unmerge_cells('A7:C7')
            wb.save(path)
            wb.close()
        self.assert_rejected(mutate)

    def test_unit_price_column_outside_details_is_not_exempt(self):
        wb = load_workbook(self.path)
        wb['领用']['F8'] = 2.286666666666667
        wb.save(self.path)
        wb.close()
        self.assert_rejected(self.change_cell('F8', 2.28666666666667))

    def test_layout_read_is_covered_by_the_source_hash(self):
        from invoice_desktop.layout import layout_plan
        def read_then_change(path):
            plan = layout_plan(path)
            self.change_cell('B4', '外部修改')(path)
            return plan
        with patch('invoice_desktop.layout.layout_plan', side_effect=read_then_change), \
                patch.object(ExcelAdapter, '_run', return_value='') as native:
            with self.assertRaisesRegex(ValueError, '工作簿在准备排版时发生变化'):
                ExcelAdapter(backend='WPS').apply(Path('shared-fonts.ps1'), self.path)
            native.assert_not_called()
        wb = load_workbook(self.path)
        self.assertEqual(wb['领用']['B4'].value, '外部修改')
        wb.close()


if __name__ == '__main__':
    unittest.main()
