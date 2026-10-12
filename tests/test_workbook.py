import datetime
import hashlib
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from tests.fixtures import template, payload

try:
    from invoice_desktop.workbook import plan_workbook, write_workbook
except ImportError:
    plan_workbook = write_workbook = None


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(plan_workbook, "需要实现不会保存输入的工作簿计划")
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = template(Path(self.tmp.name) / "表格.xlsx")

    def test_dry_run_leaves_input_bytes_unchanged(self):
        before = self.path.read_bytes()
        plan = plan_workbook(self.path, payload())
        self.assertEqual(plan.rows_to_add, 1)
        self.assertEqual(plan.blank_rows_to_remove, {"验收": 2, "领用": 2})
        self.assertEqual(self.path.read_bytes(), before)
        plan.workbook.close()

    def test_duplicate_details_and_blank_quantity_preserved_independently(self):
        one = payload()["records"][0]["rows"][0]
        one["quantity"] = ""
        two = dict(one, quantity="2", amount="5.66")
        plan = plan_workbook(self.path, payload([one, two]), datetime.date(2026, 10, 12))
        write_workbook(plan)
        wb = load_workbook(self.path)
        self.addCleanup(wb.close)
        self.assertEqual([wb["验收"].cell(r, 2).value for r in (4,5)], ["测试器材", "测试器材"])
        self.assertEqual(wb["验收"]["E4"].value, 1)
        self.assertEqual(wb["验收"]["F4"].value, "=G4/E4")
        self.assertEqual(wb["领用"]["F5"].value, 2.83)
        self.assertIsNone(wb["领用"]["H4"].value)
        self.assertEqual(wb["验收"]["H7"].value, "2026.10.12")
        self.assertEqual(wb["验收"]["G6"].value, "=SUM(G4:G5)")

    def test_expansion_and_compaction_preserve_footer_height_merges_print_area(self):
        template(self.path, blank_rows=1)
        one = payload()["records"][0]["rows"][0]
        rows = [dict(one, spec=str(i)) for i in range(5)]
        write_workbook(plan_workbook(self.path, payload(rows)))
        wb = load_workbook(self.path)
        self.addCleanup(wb.close)
        for ws in wb:
            self.assertEqual(ws["A9"].value, "合计")
            self.assertIn("A9:C9", map(str, ws.merged_cells.ranges))
            self.assertEqual([ws.row_dimensions[r].height for r in (9,10,11)], [28,32,44])
            self.assertIn("$I$11", str(ws.print_area))
            self.assertEqual(ws.row_dimensions[4].height, 19)
            self.assertEqual(ws.row_dimensions[8].height, 19)

    def test_existing_amount_conflict_does_not_overwrite(self):
        write_workbook(plan_workbook(self.path, payload()))
        changed = payload()
        changed["records"][0]["rows"][0]["amount"] = "99.00"
        before = self.path.read_bytes()
        with self.assertRaisesRegex(ValueError, "冲突"):
            plan_workbook(self.path, changed)
        self.assertEqual(self.path.read_bytes(), before)

    def test_repeated_run_is_no_change_and_keeps_remarks(self):
        write_workbook(plan_workbook(self.path, payload()))
        wb = load_workbook(self.path)
        wb["验收"]["I4"] = "保留备注"
        wb.save(self.path)
        wb.close()
        plan = plan_workbook(self.path, payload())
        self.assertEqual(plan.rows_to_add, 0)
        self.assertEqual(plan.blank_rows_to_remove, {"验收": 0, "领用": 0})
        self.assertEqual(plan.workbook["验收"]["I4"].value, "保留备注")
        plan.workbook.close()

    def test_bad_footer_date_and_quantity_are_blocked(self):
        wb = load_workbook(self.path)
        wb["验收"]["G8"] = "错误标签"
        wb.save(self.path)
        wb.close()
        with self.assertRaisesRegex(ValueError, "日期"):
            plan_workbook(self.path, payload())
        template(self.path)
        data = payload()
        data["records"][0]["rows"][0]["quantity"] = "0"
        with self.assertRaisesRegex(ValueError, "数量"):
            plan_workbook(self.path, data)

    def test_duplicate_footer_is_blocked(self):
        wb = load_workbook(self.path)
        wb["领用"]["A11"] = "合计"
        wb.save(self.path)
        wb.close()
        with self.assertRaisesRegex(ValueError, "多个"):
            plan_workbook(self.path, payload())

    def test_missing_issue_row_is_synchronized(self):
        write_workbook(plan_workbook(self.path, payload()))
        wb = load_workbook(self.path)
        for cell in wb["领用"][4]:
            cell.value = None
        wb.save(self.path)
        wb.close()
        plan = plan_workbook(self.path, payload())
        self.assertEqual(plan.rows_to_add, 1)
        write_workbook(plan)
        wb = load_workbook(self.path)
        self.addCleanup(wb.close)
        self.assertEqual(wb["领用"]["B4"].value, "测试器材")

    def test_partial_duplicate_invoice_restores_source_order_and_moves_remark(self):
        first = payload()["records"][0]["rows"][0]
        second_invoice = dict(first, sequence=2, product_name="第二发票")
        write_workbook(plan_workbook(self.path, payload([first, second_invoice])))
        wb = load_workbook(self.path)
        wb["验收"]["I5"] = "第二票备注"
        wb.save(self.path)
        wb.close()
        duplicate = dict(first, amount="5.66")
        write_workbook(plan_workbook(self.path, payload([first, duplicate, second_invoice])))
        wb = load_workbook(self.path)
        self.addCleanup(wb.close)
        self.assertEqual([wb["验收"].cell(r,1).value for r in (4,5,6)], [1,1,2])
        self.assertEqual(wb["验收"]["I6"].value, "第二票备注")
        self.assertIsNone(wb["验收"]["I5"].value)

    def test_product_text_starting_equal_sign_is_literal(self):
        data = payload()
        data["records"][0]["rows"][0]["product_name"] = "=测试器材"
        write_workbook(plan_workbook(self.path, data))
        wb = load_workbook(self.path)
        self.addCleanup(wb.close)
        self.assertEqual(wb["验收"]["B4"].data_type, "s")

    def test_moved_remark_formulas_keep_relative_absolute_and_sheet_dependencies(self):
        first = payload()["records"][0]["rows"][0]
        second = dict(first, sequence=2, product_name="第二发票", amount="20.00")
        write_workbook(plan_workbook(self.path, payload([first, second])))
        wb = load_workbook(self.path)
        wb["验收"]["I5"] = "=G5+$G$5+'领用'!G5+G4"
        wb.save(self.path)
        wb.close()
        before = self.path.read_bytes()
        plan = plan_workbook(self.path, payload([first, dict(first, amount="5.66"), second]))
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(plan.workbook["验收"]["I6"].value, "=G6+$G$6+'领用'!G6+G4")
        write_workbook(plan)

    def test_footer_formula_reference_moves_with_compaction(self):
        wb = load_workbook(self.path)
        wb["验收"]["I8"] = "=G7"
        wb.save(self.path)
        wb.close()
        plan = plan_workbook(self.path, payload())
        self.assertEqual(plan.workbook["验收"]["I6"].value, "=G5")
        plan.workbook.close()


if __name__ == "__main__":
    unittest.main()
