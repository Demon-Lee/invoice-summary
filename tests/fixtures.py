from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Border, Side


def template(path, blank_rows=3):
    wb = Workbook()
    wb.remove(wb.active)
    footer = 4 + blank_rows
    for name, first, second in (("验收", "经办人", "验收人"), ("领用", "保管人", "领用人")):
        ws = wb.create_sheet(name)
        headers = ["序号", "产品名称", "规格型号", "单位", "数量", "单价", "金额"]
        if name == "验收":
            headers.append("生产厂家（销售单位）")
        for col, header in enumerate(headers, 1):
            ws.cell(3, col, header)
        for row in range(4, footer):
            ws.row_dimensions[row].height = 19
            for col in range(1, len(headers) + 1):
                ws.cell(row, col).border = Border(bottom=Side(style="thin"))
        ws.cell(footer, 1, "合计")
        ws.cell(footer, 5, f"=SUM(E4:E{footer-1})")
        ws.cell(footer, 7, f"=SUM(G4:G{footer-1})")
        ws.merge_cells(start_row=footer, start_column=1, end_row=footer, end_column=3)
        ws.cell(footer+1, 1, first + "：")
        ws.cell(footer+1, 4, second + "：")
        ws.cell(footer+1, 7, "验收日期：" if name == "验收" else "领用日期：")
        ws.cell(footer+1, 8, "2026.1.1")
        ws.cell(footer+2, 1, "说明：保留原模板说明")
        for offset, height in enumerate((28, 32, 44)):
            ws.row_dimensions[footer+offset].height = height
        ws.print_area = f"A1:I{footer+2}"
    wb.save(path)
    wb.close()
    return Path(path)


def payload(rows=None):
    rows = rows or [dict(sequence=1, product_name="测试器材", spec="A", unit="个", quantity="1", amount="12.34", manufacturer="测试供应商")]
    return {"records": [{"sequence": 1, "pdf": "original.pdf", "grand_total": "12.34", "detail_total": "12.34", "rows": rows}], "exceptions": [], "rename_plan": []}

