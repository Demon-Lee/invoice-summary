"""Software-owned export of saved worksheets, independent of invoice rules."""
import json
import os
import shutil
import stat
import tempfile
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook

from .controller import ExcelAdapter
from .updates import sha256
from .layout import sheet_layout
from .pdf_layout import layout_reports, validate_pdf_layout
from .signatures import SLOTS
from .pdf_signatures import add_signatures

OUTPUTS = {"验收": "验收单.pdf", "领用": "领用单.pdf"}


class ProcessedWorkbookChanged(ValueError):
    """The saved workbook no longer matches the completed processing run."""


def _reparse(path):
    return bool(path.lstat().st_file_attributes & stat.FILE_ATTRIBUTE_REPARSE_POINT) if os.name == "nt" else path.is_symlink()


def _fingerprint(path):
    if not os.path.lexists(path):
        return None
    if _reparse(path) or not path.is_file():
        raise ValueError(f"输出位置不是普通文件，不能替换：{path}")
    return sha256(path)


def output_fingerprints(directory):
    directory = Path(directory).resolve()
    if directory.exists() and not directory.is_dir():
        raise ValueError(f"导出位置不是文件夹：{directory}")
    return {name: _fingerprint(directory / name) for name in OUTPUTS.values()}


def _sheet_plan(ws, filename, with_signatures=False):
    plan = sheet_layout(ws)
    plan["file_name"] = filename
    footer = plan["footer_row"]
    if with_signatures:
        row_height = ws.row_dimensions[footer + 1].height or ws.sheet_format.defaultRowHeight or 15.0
        plan["signature_slots"] = []
        for column, (key, sheet_name, label) in zip((1, 4), (slot for slot in SLOTS if slot[1] == ws.title)):
            cell = ws.cell(footer + 1, column)
            plan["signature_slots"].append({"key": key, "label": label, "row_height": float(row_height),
                                            "font_size": float(cell.font.sz or 11),
                                            "vertical_alignment": cell.alignment.vertical or "bottom"})
        plan["signature_date_label"] = "验收日期" if ws.title == "验收" else "领用日期"
    return plan


def _source_cells(ws, sheet):
    """Keep independent raw text/type/merge expectations out of native JSON."""
    merged_ends = {ws.cell(m.min_row, m.min_col).coordinate: m.max_row for m in ws.merged_cells.ranges}
    result = {}
    for address in sheet["content_cells"]:
        cell = ws[address]
        kind = "formula" if cell.data_type == "f" else "literal" if cell.data_type in ("s", "inlineStr") else "formatted"
        result[address] = {"kind": kind, "text": str(cell.value), "end_row": merged_ends.get(address, cell.row)}
    return result


def _print_plan(workbook, with_signatures=False, *, source_cells=None):
    wb = load_workbook(workbook, read_only=False, data_only=False, keep_links=False)
    try:
        missing = set(OUTPUTS) - set(wb.sheetnames)
        if missing:
            raise ValueError("工作簿缺少工作表：" + "、".join(sorted(missing)))
        sheets = [_sheet_plan(wb[name], filename, with_signatures) for name, filename in OUTPUTS.items()]
        if source_cells is not None:
            source_cells.update({sheet["name"]: _source_cells(wb[sheet["name"]], sheet) for sheet in sheets})
        return {"schema_version": 1, "sheets": sheets}
    finally:
        wb.close()


def _cleanup_owned(path, parent, prefix):
    """Only remove our own exact directory; never traverse reparse points."""
    if path is None or not path.exists():
        return
    if path.resolve() != path or path.parent != parent.resolve() or not path.name.startswith(prefix) or _reparse(path):
        raise ValueError(f"临时目录范围变化，已停止清理：{path}")
    for base, directories, files in os.walk(path, followlinks=False):
        for name in directories + files:
            if _reparse(Path(base) / name):
                raise ValueError(f"临时目录出现链接，已停止清理：{path}")
    shutil.rmtree(path)


def _validate_pdfs(stage, plan, report, source_cells):
    layouts = layout_reports(plan, report)
    files = []
    for sheet in plan["sheets"]:
        path = stage / sheet["file_name"]
        if not path.is_file() or path.stat().st_size == 0 or _reparse(path):
            raise ValueError(f"{sheet['name']} 未生成有效 PDF")
        with path.open("rb") as stream:
            if not stream.read(8).startswith(b"%PDF-"):
                raise ValueError(f"{sheet['name']} 的输出不是 PDF")
        with pdfplumber.open(path) as pdf:
            if not pdf.pages:
                raise ValueError(f"{sheet['name']} 的 PDF 没有页面")
            for page in pdf.pages:
                if abs(page.width - 841.89) > 2 or abs(page.height - 595.28) > 2:
                    width_mm, height_mm = page.width * 25.4 / 72, page.height * 25.4 / 72
                    raise ValueError(f"{sheet['name']} 第 {page.page_number} 页实际尺寸为 {width_mm:.1f} × {height_mm:.1f} 毫米，"
                                     "需要 A4 横向（297 × 210 毫米），已停止保存 PDF。"
                                     "可尝试启用 Windows 的 Microsoft Print to PDF 后重新导出。")
            validate_pdf_layout(pdf, sheet, layouts[sheet["name"]], source_cells.get(sheet["name"]))
            files.append({"sheet": sheet["name"], "name": path.name, "pages": len(pdf.pages)})
    return files


def _publish_pair(stage, directory, approved, installed):
    """Back up both outputs, then replace; roll back any completed replacement."""
    backups = {}
    if output_fingerprints(directory) != approved:
        raise ValueError("导出目录的同名文件已变化，请重新导出并确认覆盖")
    for name, digest in approved.items():
        if digest is not None:
            backup = stage / (name + ".previous")
            shutil.copyfile(directory / name, backup)
            if sha256(backup) != digest or _fingerprint(directory / name) != digest:
                raise ValueError(f"备份期间文件已变化，请重新导出：{name}")
            backups[name] = backup
    try:
        for name in OUTPUTS.values():
            if _fingerprint(directory / name) != approved[name]:
                raise ValueError(f"导出文件已变化，请重新确认：{name}")
            digest = sha256(stage / name)
            os.replace(stage / name, directory / name)
            installed[name] = digest
    except Exception as exc:
        recovery = []
        for name, digest in reversed(list(installed.items())):
            try:
                target = directory / name
                if _fingerprint(target) != digest:
                    raise ValueError("输出文件又被其他程序修改，不能覆盖")
                if name in backups:
                    os.replace(backups[name], target)
                else:
                    target.unlink()
                del installed[name]
            except Exception as rollback_error:
                recovery.append(f"{name}：{rollback_error}")
        if recovery:
            # Caller keeps stage (and any backups) until the user recovers them.
            raise ValueError(f"发布未完成：{exc}\n恢复未完成：{'；'.join(recovery)}\n请保留并核对恢复目录：{stage}") from exc
        raise ValueError(f"发布未完成，已恢复原有 PDF（如有）：{exc}") from exc


def export_workbook_pdfs(workbook, directory, approved, *, progress=None, excel=None, expected_workbook_sha256=None,
                         signatures=None):
    progress = progress or (lambda message: None)
    workbook, directory = Path(workbook).resolve(), Path(directory).resolve()
    result = {"status": "failed", "workbook": str(workbook), "output_dir": str(directory), "files": [], "cleanup": "完成"}
    scratch = stage = None
    installed = {}
    published = False
    signing_report = None
    temp_parent = Path(tempfile.gettempdir()).resolve()
    try:
        if workbook.suffix.lower() != ".xlsx" or not workbook.is_file():
            if expected_workbook_sha256 is not None:
                raise ProcessedWorkbookChanged("本次处理完成的工作簿已不可用，请重新预览并执行处理后再导出")
            raise ValueError("请选择已有的 .xlsx 工作簿")
        if not isinstance(approved, dict) or set(approved) != set(OUTPUTS.values()):
            raise ValueError("导出文件确认信息不完整，请重新导出")
        if output_fingerprints(directory) != approved:
            raise ValueError("导出目录的同名文件已变化，请重新确认")
        progress("正在读取已保存的工作簿，准备 PDF 排版…")
        scratch = Path(tempfile.mkdtemp(prefix="invoice-pdf-input-", dir=temp_parent)).resolve()
        snapshot = scratch / "workbook.xlsx"
        before = sha256(workbook)
        if expected_workbook_sha256 is not None and before != expected_workbook_sha256:
            raise ProcessedWorkbookChanged("工作簿在本次执行完成后发生变化，请重新预览并执行处理后再导出")
        shutil.copyfile(workbook, snapshot)
        if sha256(snapshot) != before or sha256(workbook) != before:
            if expected_workbook_sha256 is not None:
                raise ProcessedWorkbookChanged("工作簿在读取期间发生变化，请重新预览并执行处理后再导出")
            raise ValueError("工作簿在读取期间发生变化，请保存并关闭后重新导出")
        source_cells = {}
        plan = _print_plan(snapshot, with_signatures=signatures is not None, source_cells=source_cells)
        directory.mkdir(parents=True, exist_ok=True)
        stage = Path(tempfile.mkdtemp(prefix=".invoice-pdf-stage-", dir=directory)).resolve()
        plan_file, report_file = scratch / "plan.json", scratch / "report.json"
        plan_file.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
        progress("正在通过表格软件导出验收单和领用单，请稍候…")
        (excel or ExcelAdapter()).export_pdf(snapshot, plan_file, stage, report_file)
        report = json.loads(report_file.read_text(encoding="utf-8-sig"))
        if report.get("status") != "completed" or report.get("sheets") != plan["sheets"]:
            raise ValueError("表格软件导出结果不完整，未替换已有 PDF")
        # Always prove the unsigned PDF before signature insertion can run.
        files = _validate_pdfs(stage, plan, report, source_cells)
        if signatures is not None:
            progress("正在定位四处签字栏并添加已保存签名…")
            signing_report = add_signatures(stage, plan, signatures)
        progress("正在检查两份 PDF 并保存到导出目录…")
        if signatures is not None:
            files = _validate_pdfs(stage, plan, report, source_cells)
        if sha256(workbook) != before:
            raise ProcessedWorkbookChanged("工作簿在导出期间发生变化，请重新预览并执行处理后再导出")
        _publish_pair(stage, directory, approved, installed)
        published = True
        result.update(status="completed", files=[dict(item, path=str(directory / item["name"])) for item in files])
        if signing_report is not None:
            result["signatures"] = signing_report
    except Exception as exc:
        result["error"] = str(exc)
        missing_source = (expected_workbook_sha256 is not None and isinstance(exc, FileNotFoundError)
                          and str(exc.filename).casefold() == str(workbook).casefold())
        if isinstance(exc, ProcessedWorkbookChanged) or missing_source:
            result["requires_processing"] = True
            if missing_source:
                result["error"] = "本次处理完成的工作簿已不存在，请重新预览并执行处理后再导出"
    finally:
        cleanup_errors = []
        if installed and not published:
            cleanup_errors.append(f"保留恢复目录：{stage}")
        elif stage is not None:
            try:
                _cleanup_owned(stage, directory, ".invoice-pdf-stage-")
            except Exception as exc:
                cleanup_errors.append(str(exc))
        try:
            _cleanup_owned(scratch, temp_parent, "invoice-pdf-input-")
        except Exception as exc:
            cleanup_errors.append(str(exc))
        if cleanup_errors:
            result["cleanup"] = "；".join(cleanup_errors)
    return result
