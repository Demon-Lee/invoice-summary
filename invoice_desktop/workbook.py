"""Template-preserving workbook planning and staged in-place writes."""
import datetime
import math
import os
import re
import tempfile
import warnings
from copy import copy
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formula import Tokenizer
from openpyxl.styles import Alignment
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.print_settings import PrintArea
from openpyxl.workbook.properties import CalcProperties
from .template_rules import minimum_retained_row

HEADERS = ("序号", "产品名称", "规格型号", "单位", "数量", "单价", "金额")
SIGNERS = {"验收": ("经办人", "验收人"), "领用": ("保管人", "领用人")}
SAFE_FORMULA_FUNCTIONS = {"SUM", "MIN", "MAX", "AVERAGE", "COUNT", "COUNTA", "PRODUCT", "ROUND", "ROUNDUP", "ROUNDDOWN", "ABS", "IF", "IFERROR", "AND", "OR", "NOT", "TEXT"}


def text(value):
    return "" if value is None else str(value).strip()


def number(value, label, cents=False):
    if isinstance(value, bool):
        raise ValueError(f"{label}不是有效数字")
    try:
        result = Decimal(text(value).replace(",", ""))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{label}不是有效数字") from exc
    if not result.is_finite() or (cents and result != result.quantize(Decimal("0.01"))):
        raise ValueError(f"{label}必须为有效的逐分金额")
    return result


def _excel_number(value, label, exact=False):
    try:
        result = float(value)
        if not math.isfinite(result) or (value != 0 and result == 0):
            raise ValueError(f"{label}超出 Excel 可用数字范围")
        if exact and (Decimal(str(result)) != value or len("".join(map(str, value.as_tuple().digits)).rstrip("0")) > 15):
            raise ValueError(f"{label}超出 Excel 可准确保存的数字精度（最多 15 位有效数字）")
        return result
    except (OverflowError, InvalidOperation) as exc:
        raise ValueError(f"{label}超出 Excel 可用数字范围") from exc


def inspect_sheet(ws):
    required = HEADERS + (("生产厂家",) if ws.title == "验收" else ())
    matches = []
    for row in ws.iter_rows():
        columns = {}
        for cell in row:
            value = text(cell.value)
            key = "生产厂家" if value.startswith("生产厂家") else value
            if key in required:
                if key in columns:
                    raise ValueError(f"{ws.title} 第 {cell.row} 行表头重复：{key}")
                columns[key] = cell.column
        if all(key in columns for key in required):
            matches.append((row[0].row, columns))
    if len(matches) != 1:
        raise ValueError(f"{ws.title} 需要唯一完整表头：{'、'.join(required)}")
    header, columns = matches[0]
    footers = [r for r in range(header + 1, ws.max_row + 1) if text(ws.cell(r, 1).value).startswith("合计")]
    if len(footers) != 1:
        raise ValueError(f"{ws.title} {'检测到多个合计页脚：'+str(footers) if footers else '缺少合计页脚'}")
    footer = footers[0]
    first, second = SIGNERS[ws.title]
    if not text(ws.cell(footer+1, 1).value).startswith(first) or not text(ws.cell(footer+1, 4).value).startswith(second):
        raise ValueError(f"{ws.title} 第 {footer+1} 行签字行需要 A列{first}、D列{second}")
    if not text(ws.cell(footer+2, 1).value).startswith("说明"):
        raise ValueError(f"{ws.title} 第 {footer+2} 行缺少说明")
    date_label = "验收日期：" if ws.title == "验收" else "领用日期："
    dates = [c.column for c in ws[footer+1] if text(c.value) == date_label]
    if len(dates) != 1 or dates[0] >= ws.max_column or isinstance(ws.cell(footer+1, dates[0]+1), MergedCell):
        raise ValueError(f"{ws.title} 第 {footer+1} 行需要唯一的{date_label}及右侧日期单元格")
    for merge in ws.merged_cells.ranges:
        if merge.min_row <= footer-1 and merge.max_row > header:
            raise ValueError(f"{ws.title} 数据区域存在合并单元格：{merge}，需要处理模板")
    return header, columns, footer, dates[0]+1


def _rows(payload):
    result = []
    for record in payload.get("records", []):
        for raw in record.get("rows", []):
            row = {key: raw.get(key) for key in ("sequence", "product_name", "spec", "unit", "quantity", "amount", "manufacturer")}
            sequence = number(row["sequence"], "序号")
            if sequence <= 0 or sequence != int(sequence) or not text(row["product_name"]):
                raise ValueError("序号或产品名称缺失/无效")
            row["sequence"] = int(sequence)
            row["quantity"] = number(row["quantity"] if text(row["quantity"]) else 1, "数量")
            if row["quantity"] == 0:
                raise ValueError("数量不能为零")
            row["amount"] = number(row["amount"], "金额", cents=True)
            _excel_number(row["quantity"], "数量", exact=True)
            _excel_number(row["amount"], "金额", exact=True)
            _excel_number(row["amount"] / row["quantity"], "单价")
            for key in ("product_name", "spec", "unit", "manufacturer"):
                row[key] = text(row[key])
            if not row["manufacturer"]:
                raise ValueError("生产厂家/销方名称缺失")
            result.append(row)
    return result


def _existing(ws, columns, row_number):
    get = lambda key: ws.cell(row_number, columns[key]).value
    sequence = number(get("序号"), f"{ws.title}第{row_number}行序号")
    if sequence <= 0 or sequence != int(sequence):
        raise ValueError(f"{ws.title} 第 {row_number} 行序号无效")
    result = dict(sequence=int(sequence),
                  product_name=text(get("产品名称")), spec=text(get("规格型号")), unit=text(get("单位")),
                  quantity=number(get("数量"), "数量"),
                  amount=number(get("金额"), "金额", cents=True))
    if "生产厂家" in columns:
        result["manufacturer"] = text(get("生产厂家"))
    return result


def _detail_last_column(ws, header, columns):
    """Include all business headers, including brand and remarks, in the refill."""
    last = max([*columns.values(), *(c.column for c in ws[header] if text(c.value))])
    return max([last, *(rng.max_col for rng in ws.merged_cells.ranges
                       if rng.min_row <= header <= rng.max_row and rng.min_col <= last)])


def _populated(ws, row, last_column=None):
    cells = ws.iter_rows(min_row=row, max_row=row, max_col=last_column)
    return any(text(c.value) or c.comment is not None or c.hyperlink is not None for c in next(cells))


def _fixed_margin(sheet, row, column, areas):
    header, last_column = areas.get(sheet, (row, column))
    return row > header and column > last_column


def _capture_margin(ws, header, last_column):
    if last_column >= ws.max_column:
        return {}
    return {(c.row, c.column): copy(c)
            for row in ws.iter_rows(min_row=header+1, min_col=last_column+1)
            for c in row if not isinstance(c, MergedCell)}


def _restore_margin(ws, header, last_column, margin, original_print_ranges, original_footer, footer_delta):
    """A side legend stays at its original coordinates while the main table moves."""
    if margin:
        for row in ws.iter_rows(min_row=header+1, min_col=last_column+1):
            for cell in row:
                cell.value = cell.comment = cell.hyperlink = None
                cell._style = None
        for (row, column), saved in margin.items():
            cell = ws.cell(row, column)
            cell.value, cell.data_type = copy(saved.value), saved.data_type
            cell._style = copy(saved._style)
            cell.comment, cell.hyperlink = copy(saved.comment), copy(saved.hyperlink)
    if ws.print_area:
        ranges = [copy(rng) for rng in PrintArea.from_string(ws.print_area).ranges]
        origins = {}
        for original in original_print_ranges:
            shifted = copy(original)
            if shifted.min_row >= original_footer:
                shifted.min_row += footer_delta
            if shifted.max_row >= original_footer:
                shifted.max_row += footer_delta
            origins.setdefault(shifted.coord, []).append(original)
        for rng in ranges:
            originals = origins.get(rng.coord, [])
            included = [row for (row, col), c in margin.items()
                        if c.value is not None and rng.min_col <= col <= rng.max_col
                        and any(old.min_row <= row <= old.max_row for old in originals)]
            if included:
                rng.min_row, rng.max_row = min(rng.min_row, min(included)), max(rng.max_row, max(included))
        ws.print_area = [rng.coord for rng in ranges]


def _same_margin_cell(actual, expected):
    def link_value(link):
        return None if link is None else (link.ref, link.target, link.location, link.tooltip, link.display)
    def comment_value(comment):
        return None if comment is None else (comment.text, comment.author)
    same_type = actual.data_type == expected.data_type or actual.value is None and expected.value is None
    same_style = actual._style == expected._style or not actual.has_style and not expected.has_style
    return (actual.value == expected.value and same_type and same_style
            and comment_value(actual.comment) == comment_value(expected.comment)
            and link_value(actual.hyperlink) == link_value(expected.hyperlink))


def _matches_sheet(ws, header, columns, footer, rows, last_column=None, minimum_row=None):
    """Recognize an already-written pending refill, without authorizing normal skips."""
    last_column = last_column or _detail_last_column(ws, header, columns)
    if footer != max(header + len(rows), minimum_row or header) + 1:
        return False
    for r, expected in enumerate(rows, header + 1):
        for label in ("序号", "数量", "金额"):
            cell = ws.cell(r, columns[label])
            if cell.data_type != "n" or cell.value is None or isinstance(cell.value, bool):
                return False
        for label in ("产品名称", "规格型号", "单位", "生产厂家"):
            if label in columns:
                cell = ws.cell(r, columns[label])
                if cell.value is not None and cell.data_type != "s":
                    return False
        try:
            existing = _existing(ws, columns, r)
            if any(existing[key] != expected[key] for key in existing):
                return False
            price = ws.cell(r, columns["单价"])
            if ws.title == "验收":
                if price.data_type != "f" or price.value != f"={get_column_letter(columns['金额'])}{r}/{get_column_letter(columns['数量'])}{r}":
                    return False
            elif price.data_type != "n" or price.value is None or isinstance(price.value, bool) or not math.isclose(float(number(price.value, "单价")), _excel_number(expected["amount"] / expected["quantity"], "单价"), rel_tol=1e-12, abs_tol=0):
                return False
        except (ValueError, TypeError, OverflowError, InvalidOperation):
            return False
        for cell in ws[r]:
            if cell.column > last_column:
                continue
            if cell.comment is not None or cell.hyperlink is not None:
                return False
            if cell.column not in columns.values() and cell.value is not None:
                return False
    for row in ws.iter_rows(min_row=header+len(rows)+1, max_row=footer-1, max_col=last_column):
        if any(c.value is not None or c.comment is not None or c.hyperlink is not None for c in row):
            return False
    for label in ("数量", "金额"):
        letter = get_column_letter(columns[label])
        cell = ws.cell(footer, columns[label])
        if cell.data_type != "f" or cell.value != f"=SUM({letter}{header+1}:{letter}{footer-1})":
            return False
    return True


def _shift_rows(ws, start, delta, fixed_area=None):
    """openpyxl does not shift merge/dimension/print metadata by itself."""
    if not delta:
        return
    unsupported = []
    for present, label in ((ws.tables, "Excel 表格"), (ws.auto_filter.ref, "自动筛选区域"),
                           (ws.conditional_formatting, "条件格式"),
                           (ws.data_validations and ws.data_validations.dataValidation, "数据验证"),
                           (ws._charts, "图表"), (ws._images, "图片"), (ws._pivots, "透视表"),
                           (ws.row_breaks.brk, "手动分页")):
        if present:
            unsupported.append(label)
    if unsupported:
        raise ValueError(f"{ws.title} 需要增减明细行，但包含无法安全随行移动的结构：{'、'.join(unsupported)}，请调整模板")
    affected = min(start, start + delta)
    if ws.print_title_rows:
        title_rows = str(ws.print_title_rows).rsplit("!", 1)[-1]
        match = re.fullmatch(r"\$?([1-9][0-9]*):\$?([1-9][0-9]*)", title_rows)
        if match is None or int(match[2]) >= affected:
            raise ValueError(f"{ws.title} 重复打印行涉及待移动区域，请调整模板")
    if ws.freeze_panes and CellRange(str(ws.freeze_panes)).max_row >= affected:
        raise ValueError(f"{ws.title} 冻结窗格涉及待移动区域，请调整模板")
    if any(cell.hyperlink is not None
           and not (fixed_area and cell.row > fixed_area[0] and cell.column > fixed_area[1])
           for row in ws.iter_rows(min_row=affected) for cell in row):
        raise ValueError(f"{ws.title} 超链接涉及待移动区域，请调整模板")
    merges = [copy(rng) for rng in ws.merged_cells.ranges if rng.min_row >= start]
    if any(rng.min_row < start <= rng.max_row for rng in ws.merged_cells.ranges):
        raise ValueError(f"{ws.title} 合并单元格跨越移动区域")
    dimensions = {r: copy(dim) for r, dim in ws.row_dimensions.items() if r >= start}
    print_ranges = [copy(rng) for rng in PrintArea.from_string(ws.print_area).ranges] if ws.print_area else []
    for rng in merges:
        ws.unmerge_cells(str(rng))
    if delta > 0:
        ws.insert_rows(start, delta)
    else:
        ws.delete_rows(start + delta, -delta)
    for r in list(ws.row_dimensions):
        if r >= min(start, start+delta):
            del ws.row_dimensions[r]
    for r, dim in dimensions.items():
        dim.index = r + delta
        ws.row_dimensions[r+delta] = dim
    for rng in merges:
        rng.shift(row_shift=delta)
        ws.merge_cells(str(rng))
    for rng in print_ranges:
        if rng.min_row >= start:
            rng.min_row += delta
        if rng.max_row >= start:
            rng.max_row += delta
    if print_ranges:
        ws.print_area = [rng.coord for rng in print_ranges]


def _remap_formula(value, sheet, row_maps, fixed_areas=None):
    """Track references to moved cells, including absolute and cross-sheet refs."""
    tokens = Tokenizer(value).items
    for token in tokens:
        if token.type == "FUNC" and token.subtype == "OPEN" and token.value[:-1].upper() not in SAFE_FORMULA_FUNCTIONS:
            raise ValueError(f"{sheet} 公式函数无法安全随行移动：{value}")
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        prefix, reference = token.value.rsplit("!", 1) if "!" in token.value else ("", token.value)
        target_sheet = prefix.strip("'").replace("''", "'") if prefix else sheet
        if prefix and ("[" in prefix or ":" in prefix):
            raise ValueError(f"{sheet} 外部或跨表区域公式无法安全随行移动：{value}")
        parts = reference.split(":")
        fixed_area = (fixed_areas or {}).get(target_sheet)
        if fixed_area and len(parts) == 2 and all(re.fullmatch(r"\$?[A-Za-z]{1,3}\$?[1-9][0-9]*", part) for part in parts):
            region = CellRange(reference)
            header, last_column = fixed_area
            if (region.min_col <= last_column < region.max_col and region.max_row > header
                    and any(region.min_row <= old <= region.max_row and old != new
                            for old, new in row_maps.get(target_sheet, {}).items())):
                raise ValueError(f"{sheet} 公式区域跨越主表和固定右侧文字，无法安全随行移动：{value}")
        mapped = []
        fixed_flags, changed = [], False
        for part in parts:
            match = re.fullmatch(r"(\$?[A-Za-z]{1,3}\$?)([1-9][0-9]*)", part)
            if match:
                row = int(match[2])
                column = column_index_from_string(match[1].replace("$", ""))
                fixed = _fixed_margin(target_sheet, row, column, fixed_areas or {})
                fixed_flags.append(fixed)
                moved = row if fixed else row_maps.get(target_sheet, {}).get(row, row)
                if moved is None:
                    raise ValueError(f"{sheet} 公式引用待移除的明细行，无法安全移动：{value}")
                mapped.append(match[1] + str(moved))
                changed = changed or moved != row
            elif re.fullmatch(r"\$?[A-Za-z]{1,3}", part) and len(parts) == 2:
                mapped.append(part)
            else:
                raise ValueError(f"{sheet} 命名或特殊区域公式无法安全随行移动：{value}")
        if changed and len(set(fixed_flags)) > 1:
            raise ValueError(f"{sheet} 公式区域跨越主表和固定右侧文字，无法安全随行移动：{value}")
        token.value = (prefix + "!" if prefix else "") + ":".join(mapped)
    return "=" + "".join(token.value for token in tokens)


@dataclass
class WorkbookPlan:
    path: Path
    workbook: object
    rows_to_add: int
    blank_rows_to_remove: dict
    total: Decimal
    sheet_details: dict
    rows_to_correct: int = 0
    rows_to_fill: int = 0
    workbook_matches: bool = False
    fixed_margins: dict = field(default_factory=dict)
    template_rules: object = None


def plan_workbook(path, payload, run_date=None, *, template_rules=None):
    path = Path(path).resolve()
    if path.suffix.lower() != ".xlsx" or not path.is_file():
        raise ValueError("请选择现有的 .xlsx 工作簿")
    with warnings.catch_warnings(record=True) as noticed:
        warnings.simplefilter("always")
        wb = load_workbook(path, rich_text=True)
    if noticed:
        wb.close()
        raise ValueError("工作簿包含当前写入组件无法完整保留的结构：" + "; ".join(str(n.message) for n in noticed))
    try:
        missing = set(SIGNERS) - set(wb.sheetnames)
        if missing:
            raise ValueError("缺少工作表：" + "、".join(sorted(missing)))
        rows = _rows(payload)
        if not rows:
            raise ValueError("没有可重填的发票明细，未清空工作簿")
        issue_day = run_date or datetime.date.today()
        acceptance_day = issue_day - datetime.timedelta(days=1)
        issue_date_text = f"{issue_day.year}.{issue_day.month}.{issue_day.day}"
        acceptance_date_text = f"{acceptance_day.year}.{acceptance_day.month}.{acceptance_day.day}"
        removed, details, added_counts, matching = {}, {}, [], []
        formulas = [(ws.title, c.row, c.column, c.value) for ws in wb for row in ws for c in row if c.data_type == "f"]
        row_maps, regenerated, discarded, fixed_areas, margins = {}, set(), set(), {}, {}
        for name in SIGNERS:
            ws = wb[name]
            header, columns, footer, date_col = inspect_sheet(ws)
            last_column = _detail_last_column(ws, header, columns)
            minimum_row = minimum_retained_row(ws, header, last_column, template_rules)
            if minimum_row is not None and footer <= minimum_row:
                raise ValueError(f"{name} 侧边标识结束于第 {minimum_row} 行，但合计在第 {footer} 行；请恢复完整模板，使合计位于侧边文字之后")
            last = max(header + len(rows), minimum_row or header)
            required_rows = last - header
            fixed_areas[name] = (header, last_column)
            margin = _capture_margin(ws, header, last_column)
            original_print_ranges = [copy(rng) for rng in PrintArea.from_string(ws.print_area).ranges] if ws.print_area else []
            if required_rows != footer-header-1 and any(rng.max_col > last_column and rng.max_row > header for rng in ws.merged_cells.ranges):
                raise ValueError(f"{name} 主表右侧固定区域存在合并单元格，无法安全增减明细行，请调整模板")
            original_footer, original_max_row = footer, ws.max_row
            template_row = header + 1
            if template_row >= footer:
                raise ValueError(f"{name} 缺少可继承格式的明细行，请使用含空白明细行的模板")
            template_styles = [copy(c._style) for c in ws[template_row]]
            template_height = ws.row_dimensions[template_row].height
            matching.append(_matches_sheet(ws, header, columns, footer, rows, last_column, minimum_row))
            old_count = sum(_populated(ws, r, last_column) for r in range(header+1, footer))
            added_counts.append(max(0, len(rows) - old_count))
            discarded.update((name, r, col) for sheet, r, col, value in formulas if sheet == name and header < r < original_footer and col <= last_column)
            expansion = max(0, required_rows - (footer-header-1))
            _shift_rows(ws, footer, expansion, fixed_areas[name])
            footer += expansion
            for r in range(header+1, footer):
                if r >= original_footer:
                    ws.row_dimensions[r].height = template_height
                    for col, style in enumerate(template_styles, 1):
                        ws.cell(r, col)._style = copy(style)
                for cell in next(ws.iter_rows(min_row=r, max_row=r, max_col=last_column)):
                    cell.value = None
                    cell.hyperlink = cell.comment = None
            for r, row in enumerate(rows, header+1):
                values = {"序号": row["sequence"], "产品名称": row["product_name"], "规格型号": row["spec"],
                          "单位": row["unit"], "数量": _excel_number(row["quantity"], "数量", exact=True), "金额": _excel_number(row["amount"], "金额", exact=True),
                          "生产厂家": row["manufacturer"]}
                values["单价"] = (f"={get_column_letter(columns['金额'])}{r}/{get_column_letter(columns['数量'])}{r}"
                                if name == "验收" else _excel_number(row["amount"] / row["quantity"], "单价"))
                for label, col in columns.items():
                    cell = ws.cell(r, col, values[label])
                    if label in ("产品名称", "规格型号", "单位", "生产厂家"):
                        cell.data_type = "s"
                    align = copy(cell.alignment)
                    align.horizontal = align.vertical = "center"
                    align.wrap_text = True
                    align.shrink_to_fit = False
                    cell.alignment = align
            count = footer-last-1
            removed[name] = count
            _shift_rows(ws, footer, -count, fixed_areas[name])
            footer -= count
            if expansion or count:
                _restore_margin(ws, header, last_column, margin, original_print_ranges, original_footer, footer-original_footer)
            row_maps[name] = {r: r if r <= last else None for r in range(header+1, original_footer)}
            for r in range(original_footer, original_max_row+1):
                row_maps[name][r] = r + footer - original_footer
            for label in ("数量", "金额"):
                regenerated.add((name, original_footer, columns[label]))
            regenerated.add((name, original_footer+1, date_col))
            for label in ("数量", "金额"):
                letter = get_column_letter(columns[label])
                ws.cell(footer, columns[label], f"=SUM({letter}{header+1}:{letter}{footer-1})" if footer > header+1 else "=0")
            ws.cell(footer+1, date_col, acceptance_date_text if name == "验收" else issue_date_text)
            details[name] = {"headerRow": header, "footerRow": footer, "columns": columns, "detailLastColumn": last_column, "fixedTextCellsToPreserve": sum(c.value is not None for c in margin.values()), "rowsToAdd": added_counts[-1], "rowsToCorrect": 0, "blankRowsToRemove": count, "oldRowsToClear": old_count, "rowsToFill": len(rows), "minimumRetainedRow": minimum_row, "requiredBlankRows": last-header-len(rows)}
            inspect_sheet(ws)
        moved = any(old != new for mapping in row_maps.values() for old, new in mapping.items())
        if moved:
            if wb.defined_names or any(sheet.defined_names for sheet in wb):
                raise ValueError("工作簿包含命名区域，当前无法安全随行移动，请调整模板")
            for name, row, col, formula in formulas:
                if (name, row, col) in regenerated or (name, row, col) in discarded:
                    continue
                target_row = row if _fixed_margin(name, row, col, fixed_areas) else row_maps.get(name, {}).get(row, row)
                if target_row is None or not isinstance(formula, str):
                    raise ValueError(f"{name} 公式无法安全随行移动")
                wb[name].cell(target_row, col, _remap_formula(formula, name, row_maps, fixed_areas))
        for name, (header, last_column) in fixed_areas.items():
            margins[name] = _capture_margin(wb[name], header, last_column)
        wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
        return WorkbookPlan(path, wb, max(added_counts), removed, sum((r["amount"] for r in rows), Decimal("0.00")), details, rows_to_fill=len(rows), workbook_matches=all(matching), fixed_margins=margins, template_rules=template_rules)
    except Exception:
        wb.close()
        raise


def write_workbook(plan):
    handle = tempfile.NamedTemporaryFile(prefix=".invoice-write-", suffix=".xlsx", dir=plan.path.parent, delete=False)
    staged = Path(handle.name)
    handle.close()
    try:
        plan.workbook.save(staged)
        check = load_workbook(staged, rich_text=True)
        try:
            for name in SIGNERS:
                header, columns, footer, date_col = inspect_sheet(check[name])
                last_column = _detail_last_column(check[name], header, columns)
                details = plan.sheet_details[name]
                minimum_row = minimum_retained_row(check[name], header, last_column, plan.template_rules)
                expected_last = max(header+details["rowsToFill"], details["minimumRetainedRow"] or header)
                if (header != details["headerRow"] or footer != expected_last+1
                        or minimum_row != details["minimumRetainedRow"]):
                    raise ValueError(f"{name} 写后明细保留区或页脚位置不符合计划，未替换原工作簿")
                if not _populated(check[name], header+details["rowsToFill"], last_column):
                    raise ValueError(f"{name} 写后末条明细缺失，未替换原工作簿")
                for row in check[name].iter_rows(min_row=header+details["rowsToFill"]+1, max_row=footer-1, max_col=last_column):
                    if any(c.value is not None or c.comment is not None or c.hyperlink is not None for c in row):
                        raise ValueError(f"{name} 必要空白明细行存在非空内容，未替换原工作簿")
                for label in ("数量", "金额"):
                    letter = get_column_letter(columns[label])
                    expected = f"=SUM({letter}{header+1}:{letter}{footer-1})" if footer > header+1 else "=0"
                    if check[name].cell(footer, columns[label]).value != expected:
                        raise ValueError(f"{name} 写后总计公式不符合计划")
                for (row, col), expected in (plan.fixed_margins or {}).get(name, {}).items():
                    if not _same_margin_cell(check[name].cell(row, col), expected):
                        raise ValueError(f"{name} 写后右侧模板内容未完整保留：{get_column_letter(col)}{row}，未替换原工作簿")
        finally:
            check.close()
        os.replace(staged, plan.path)
    finally:
        plan.workbook.close()
        if staged.exists():
            staged.unlink()
