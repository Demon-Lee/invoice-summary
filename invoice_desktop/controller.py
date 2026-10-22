"""Read-only preview and guarded execution of one pinned skill engine."""
import contextlib
import copy
import ctypes
import importlib.util
import io
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

from .updates import read_manifest, sha256
from .workbook import plan_workbook, write_workbook, number
from .corrections import apply_corrections, row_changes
from .template_rules import load_template_rules

WRITE_MODE = "replace-details-v1"


def resource_root():
    return Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[1]))


def application_root():
    """The directory users copy, independent of cwd and _internal resources."""
    return Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parents[1]


def load_engine(directory, name):
    path = Path(directory) / name
    spec = importlib.util.spec_from_file_location("invoice_skill_" + sha256(path)[:12], path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    with contextlib.ExitStack():
        original = sys.dont_write_bytecode
        sys.dont_write_bytecode = True
        try:
            spec.loader.exec_module(module)
        finally:
            sys.dont_write_bytecode = original
    return module


class ExcelAdapter:
    def _run(self, script, extra, *, timeout=240):
        args = ["powershell.exe", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", str(script), *extra]
        try:
            result = subprocess.run(args, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=timeout,
                                    creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        except (OSError, subprocess.TimeoutExpired) as exc:
            raise ValueError(f"Excel 操作未完成：{exc}") from exc
        if result.returncode:
            raise ValueError("Excel 不可用或操作失败：" + (result.stderr or result.stdout).strip()[-2000:])

    def check(self, workbook):
        if os.name != "nt":
            raise ValueError("此版本仅支持 Windows 与 Microsoft Excel 桌面版")
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.CreateFileW.argtypes = [ctypes.c_wchar_p, ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p, ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p]
        kernel.CreateFileW.restype = ctypes.c_void_p
        kernel.CloseHandle.argtypes = [ctypes.c_void_p]
        handle = kernel.CreateFileW(str(workbook), 0x80000000 | 0x40000000, 0, None, 3, 0, None)
        if handle == ctypes.c_void_p(-1).value:
            raise ValueError("工作簿被占用或没有写入权限，请关闭 Excel 中的该文件后重新预览")
        kernel.CloseHandle(handle)
        self._run(resource_root() / "resources/check_excel.ps1", [])

    def apply(self, script, workbook):
        from .layout import layout_plan
        from openpyxl import load_workbook

        def contents(path):
            book = load_workbook(path, data_only=False)
            try:
                return [(sheet.title,
                         tuple((cell.coordinate, cell.data_type, cell.value)
                               for cell in sorted(sheet._cells.values(), key=lambda cell: (cell.row, cell.column))
                               if cell.value is not None),
                         tuple(sorted(str(area) for area in sheet.merged_cells.ranges)))
                        for sheet in book.worksheets]
            finally:
                book.close()

        workbook = Path(workbook).resolve()
        before_hash = sha256(workbook)
        before_contents = contents(workbook)
        staged = None
        try:
            with tempfile.NamedTemporaryFile(prefix=".invoice-format-", suffix=".xlsx",
                                             dir=workbook.parent, delete=False) as handle:
                staged = Path(handle.name)
            shutil.copy2(workbook, staged)
            if sha256(staged) != before_hash:
                raise ValueError("工作簿在准备排版时发生变化，请重新预览")
            with tempfile.TemporaryDirectory(prefix="invoice-layout-") as temporary:
                plan_file = Path(temporary) / "plan.json"
                plan_file.write_text(json.dumps(layout_plan(staged), ensure_ascii=False), encoding="utf-8")
                self._run(resource_root() / "resources/run_excel.ps1", [
                    "-EngineScript", str(script), "-WorkbookInput", str(staged),
                    "-LayoutPlanFile", str(plan_file),
                ], timeout=480)
            if contents(staged) != before_contents:
                raise ValueError("排版副本的单元格内容或合并结构发生变化，已停止替换工作簿")
            if sha256(workbook) != before_hash:
                raise ValueError("工作簿在排版期间被修改，已停止替换；请重新预览")
            os.replace(staged, workbook)
        finally:
            if staged is not None and staged.exists():
                staged.unlink()

    def export_pdf(self, workbook, plan_file, output_dir, report_file):
        self._run(resource_root() / "resources/export_pdf.ps1", [
            "-WorkbookInput", str(workbook), "-PlanFile", str(plan_file),
            "-OutputDirectory", str(output_dir), "-ReportFile", str(report_file),
        ], timeout=270)


@dataclass
class Preview:
    pdf_dir: Path
    workbook: Path
    engine_dir: Path
    engine_version: str
    payload: dict
    fingerprints: dict
    issues: list = field(default_factory=list)
    rows_to_add: int = 0
    blank_rows_to_remove: dict = field(default_factory=dict)
    sheet_details: dict = field(default_factory=dict)
    total: str = "0.00"
    resume_pending: bool = False
    original_payload: dict = field(default_factory=dict)
    corrections: dict = field(default_factory=dict)
    base_issues: list = field(default_factory=list)
    rows_to_correct: int = 0
    workbook_written: bool = False
    write_phase_uncertain: bool = False
    rows_to_fill: int = 0
    workbook_matches: bool = False

    @property
    def clean(self):
        return not self.issues

    @property
    def has_changes(self):
        return self.resume_pending or self.rows_to_fill > 0 or any(e.get("status") == "ready" for e in self.payload.get("rename_plan", []))


class Controller:
    def __init__(self, engine_dir, progress=None, excel=None, state_dir=None):
        self.engine_dir = Path(engine_dir).resolve()
        self.progress = progress or (lambda message: None)
        self.excel = excel or ExcelAdapter()
        self.last_scratch = None
        self.cleanup_error = None
        self.state_dir = Path(state_dir or Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / "InvoiceSummary/runs").resolve()

    def _pending_path(self, directory, workbook):
        key = hashlib.sha256(json.dumps([str(directory).casefold(), str(workbook).casefold()]).encode()).hexdigest()
        return self.state_dir / (key + ".json")

    def _invoice_hashes(self, fingerprints, workbook):
        return sorted(digest for path, digest in fingerprints.items() if path != str(workbook))

    def _save_pending(self, preview, workbook_written=None):
        self.state_dir.mkdir(parents=True, exist_ok=True)
        path = self._pending_path(preview.pdf_dir, preview.workbook)
        staged = path.with_suffix(".staged.json")
        state = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {
            "engine_version": preview.engine_version,
            "invoice_hashes": self._invoice_hashes(preview.fingerprints, preview.workbook),
            "workbook_written": False,
            "workbook_before_sha256": preview.fingerprints[str(preview.workbook)]}
        state["write_mode"] = WRITE_MODE
        if workbook_written is not None:
            state["workbook_written"] = workbook_written
        if preview.corrections and "manual" not in state:
            destinations = {p["source"]: p["destination"] for p in preview.payload["rename_plan"]}
            state["manual"] = {
                "changes": preview.corrections,
                "workbook_before_sha256": preview.fingerprints[str(preview.workbook)],
                "invoices": [{"source": r["pdf"], "destination": destinations[r["pdf"]], "sha256": preview.fingerprints[r["pdf"]], "sequence": r["sequence"], "row_count": len(r["rows"])} for r in preview.original_payload["records"]],
            }
        try:
            staged.write_text(json.dumps(state, indent=2), encoding="utf-8")
            os.replace(staged, path)
        finally:
            if staged.exists():
                staged.unlink()
        return state

    @contextlib.contextmanager
    def _scratch(self):
        parent = Path(tempfile.gettempdir()).resolve()
        path = Path(tempfile.mkdtemp(prefix="invoice-run-", dir=parent)).resolve()
        self.last_scratch = path
        self.cleanup_error = None
        try:
            yield path
        finally:
            try:
                if path.parent != parent or not path.name.startswith("invoice-run-"):
                    raise ValueError("临时目录越界，已停止清理")
                shutil.rmtree(path)
            except Exception as exc:
                self.cleanup_error = f"{path}：{exc}"

    def _invoke(self, kind, args, report):
        name = "collect_excel_totals.py" if kind == "totals" else "invoice_pipeline.py"
        module = load_engine(self.engine_dir, name)
        output = io.StringIO()
        with contextlib.redirect_stdout(output), contextlib.redirect_stderr(output):
            try:
                status = module.main(args)
            except SystemExit as exc:
                raise ValueError(output.getvalue().strip() or f"识别脚本异常退出：{exc.code}") from exc
        if status not in (None, 0):
            raise ValueError(output.getvalue().strip() or "识别脚本执行失败")
        data = json.loads(Path(report).read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            raise ValueError("识别结果格式不兼容，请升级软件")
        return data

    def _fingerprints(self, directory, workbook):
        parser = load_engine(self.engine_dir, "invoice_pipeline.py")
        classifier = getattr(parser, "classify_pdf_document", None)
        # New rules report skipped files too: pin the directory so that list
        # cannot become stale between preview and execution. Hashing is not
        # invoice numbering, renaming, writing or amount calculation.
        files = [workbook, *(p for p in sorted(directory.glob("*.pdf"))
                             if callable(classifier) or not parser.is_verification_report_pdf(p))]
        return {str(path): sha256(path) for path in files}

    def _actual_rename_plan(self, preview):
        actual = []
        for item in preview.payload["rename_plan"]:
            source, destination = Path(item["source"]), Path(item["destination"])
            if source == destination and source.is_file():
                status = "unchanged"
            elif destination.is_file() and not source.exists():
                status = "renamed"
            else:
                status = "not_completed"
            actual.append(dict(item, status=status))
        return actual

    def preview(self, pdf_dir, workbook):
        directory, book = Path(pdf_dir).resolve(), Path(workbook).resolve()
        if not directory.is_dir():
            raise ValueError("请选择有效的 PDF 文件夹")
        if book.suffix.lower() != ".xlsx" or not book.is_file():
            raise ValueError("请选择现有的 .xlsx 工作簿")
        manifest = read_manifest(self.engine_dir)
        self.progress("正在读取发票，仅作预览…")
        before = self._fingerprints(directory, book)
        with self._scratch() as scratch:
            report = scratch / "preview.json"
            data = self._invoke("pipeline", ["preview", "--pdf-dir", str(directory), "--report", str(report)], report)
            result = Preview(directory, book, self.engine_dir, manifest["version"], data, before, original_payload=copy.deepcopy(data))
            pending = self._pending_path(directory, book)
            if pending.exists():
                result.resume_pending = True
                try:
                    state = json.loads(pending.read_text(encoding="utf-8"))
                    result.workbook_written = state.get("workbook_written", False)
                    if not isinstance(result.workbook_written, bool):
                        raise ValueError("写表阶段记录无效")
                    if state["engine_version"] != result.engine_version:
                        result.base_issues.append("上次处理未完成，请先回退到原识别版本再继续处理")
                    if state["invoice_hashes"] != self._invoice_hashes(before, book):
                        result.base_issues.append("上次处理未完成且发票内容或批次已变化，请恢复原批次后继续")
                    if state.get("write_mode") not in (None, WRITE_MODE):
                        result.base_issues.append("上次处理使用不同的写表策略，请使用原软件完成后再重填")
                    before_hash = state.get("workbook_before_sha256") or (state.get("manual") or {}).get("workbook_before_sha256")
                    if before_hash is not None and (not isinstance(before_hash, str) or len(before_hash) != 64 or any(c not in "0123456789abcdef" for c in before_hash)):
                        raise ValueError("写表前内容哈希记录无效")
                    if not result.workbook_written and before_hash is not None and before_hash != before[str(book)]:
                        result.workbook_written = True
                        result.write_phase_uncertain = True
                    if not result.base_issues and state.get("manual"):
                        self._restore_manual(result, state["manual"])
                        self._rebuild_rename(result)
                except (OSError, ValueError, KeyError, TypeError) as exc:
                    result.base_issues.append("未完成处理记录无法读取：" + str(exc))
            self._validate_preview(result)
        if self.cleanup_error:
            issue = "临时文件清理未完成：" + self.cleanup_error
            result.base_issues.append(issue)
            result.issues.append(issue)
        self.progress("预览结束：" + ("可处理" if result.clean else "存在待解决异常"))
        return result

    def _restore_manual(self, preview, saved):
        bindings = saved["invoices"]
        records = preview.original_payload.get("records", [])
        if len(bindings) != len(records) or [r["sequence"] for r in bindings] != list(range(1, len(records)+1)):
            raise ValueError("未完成人工校正的发票条数或顺序不一致")
        restored, used = [], set()
        for binding in bindings:
            paths = {str(Path(binding[key]).resolve()) for key in ("source", "destination")}
            matches = [i for i, r in enumerate(records) if i not in used and str(Path(r["pdf"]).resolve()) in paths and preview.fingerprints.get(r["pdf"]) == binding["sha256"]]
            if len(matches) != 1:
                raise ValueError("无法唯一对应未完成人工校正的原始发票")
            index = matches[0]
            used.add(index)
            record = copy.deepcopy(records[index])
            if len(record["rows"]) != binding["row_count"]:
                raise ValueError("人工校正恢复时明细条数发生变化，请先改进识别")
            record["sequence"] = binding["sequence"]
            for row in record["rows"]:
                row["sequence"] = record["sequence"]
            restored.append(record)
        preview.original_payload["records"] = restored
        preview.corrections = copy.deepcopy(saved["changes"])
        preview.payload = apply_corrections(preview.original_payload, preview.corrections)

    def _rebuild_rename(self, preview):
        with self._scratch() as scratch:
            records, report = scratch / "corrected.json", scratch / "rename-preview.json"
            records.write_text(json.dumps(preview.payload, ensure_ascii=False), encoding="utf-8")
            renamed = self._invoke("pipeline", ["rename", "--records", str(records), "--report", str(report)], report)
            preview.payload["rename_plan"] = renamed["rename_plan"]
        if self.cleanup_error:
            preview.base_issues.append("临时文件清理未完成：" + self.cleanup_error)

    def _validate_preview(self, result):
        result.issues = list(result.base_issues)
        result.rows_to_add = result.rows_to_correct = result.rows_to_fill = 0
        result.workbook_matches = False
        result.blank_rows_to_remove, result.sheet_details = {}, {}
        result.total = "0.00"
        data = result.payload
        for item in data.get("exceptions", []):
            result.issues.append(f"{Path(item.get('pdf', '')).name} · {item.get('field', '')}：{item.get('message', '')}")
        records, plan = data.get("records", []), data.get("rename_plan", [])
        if not records:
            result.issues.append("没有可处理的发票")
        if len(plan) != len(records):
            result.issues.append("发票与改名计划不对应，识别接口需要适配")
        sources, targets = set(), set()
        for item in plan:
            source, destination = Path(item["source"]).resolve(), Path(item["destination"]).resolve()
            if source.parent != result.pdf_dir or destination.parent != result.pdf_dir:
                result.issues.append("改名计划超出所选文件夹")
            if item.get("status") not in ("ready", "unchanged"):
                result.issues.append(f"文件重名或改名状态异常：{destination.name}")
            if str(source).casefold() in sources or str(destination).casefold() in targets:
                result.issues.append("改名计划包含重复来源或目标")
            sources.add(str(source).casefold())
            targets.add(str(destination).casefold())
        try:
            if [r["sequence"] for r in records] != list(range(1, len(records)+1)):
                raise ValueError("发票序号必须从 1 连续且逐票唯一")
            total = sum((number(r["grand_total"], "发票金额", cents=True) for r in records), Decimal("0.00"))
            result.total = format(total, ".2f")
            for record in records:
                grand = number(record["grand_total"], "发票金额", cents=True)
                detail = number(record["detail_total"], "明细金额", cents=True)
                row_total = sum((number(r["amount"], "明细金额", cents=True) for r in record.get("rows", [])), Decimal("0.00"))
                if not record.get("rows") or grand != detail or grand != row_total:
                    raise ValueError(f"序号 {record['sequence']} 发票金额与明细金额不一致；漏项请改进 skill，字段错误可人工校正")
                if any(r.get("sequence") != record["sequence"] for r in record["rows"]):
                    raise ValueError("同一发票的明细序号不一致")
            self.progress("正在检查工作簿表头、页脚和明细映射…")
            writer_plan = plan_workbook(result.workbook, data, template_rules=load_template_rules(self.engine_dir))
            result.rows_to_add, result.rows_to_correct = writer_plan.rows_to_add, writer_plan.rows_to_correct
            result.rows_to_fill, result.workbook_matches = writer_plan.rows_to_fill, writer_plan.workbook_matches
            result.blank_rows_to_remove, result.sheet_details = writer_plan.blank_rows_to_remove, writer_plan.sheet_details
            writer_plan.workbook.close()
            if result.write_phase_uncertain and not result.workbook_matches:
                result.issues.append("工作簿已变化但写表完成记录缺失，两表未完整对应本次重填结果；请核实并恢复原工作簿后重试")
            if len(result.sheet_details) != 2:
                result.issues.append("工作簿映射不完整")
            from openpyxl import load_workbook
            wb = load_workbook(result.workbook, read_only=True)
            try:
                for ws in wb:
                    if not any(row[0] == "序号" for row in ws.iter_rows(values_only=True)):
                        result.issues.append(f"{ws.title} 无序号表头，不符合当前 skill 字体处理约定")
            finally:
                wb.close()
        except (ValueError, KeyError, TypeError) as exc:
            result.issues.append(str(exc))
        if result.fingerprints != self._fingerprints(result.pdf_dir, result.workbook):
            result.issues.append("文件在预览过程中发生变化，请重新预览")

    def correct(self, preview, record_index, row_index, values=None):
        if preview.resume_pending:
            raise ValueError("请先完成上次处理；恢复期间固定使用已保存的人工校正")
        if preview.engine_dir != self.engine_dir or read_manifest(self.engine_dir)["version"] != preview.engine_version:
            raise ValueError("识别版本发生变化，请重新预览")
        if preview.fingerprints != self._fingerprints(preview.pdf_dir, preview.workbook):
            raise ValueError("预览后输入文件或目录发生变化，请重新预览")
        result = copy.deepcopy(preview)
        try:
            original_row = result.original_payload["records"][record_index]["rows"][row_index]
        except (KeyError, IndexError, TypeError) as exc:
            raise ValueError("请选择有效的预览明细") from exc
        key = f"{record_index}:{row_index}"
        changes = row_changes(original_row, values) if values is not None else {}
        if changes:
            result.corrections[key] = changes
        else:
            result.corrections.pop(key, None)
        result.payload = apply_corrections(result.original_payload, result.corrections)
        self.progress("正在检查人工校正后的金额、文件名和工作簿映射…")
        self._rebuild_rename(result)
        self._validate_preview(result)
        return result

    def apply(self, preview):
        if not preview.clean:
            raise ValueError("存在未解决异常，不能执行处理")
        if preview.engine_dir != self.engine_dir or read_manifest(self.engine_dir)["version"] != preview.engine_version:
            raise ValueError("识别版本发生变化，请重新预览")
        if preview.fingerprints != self._fingerprints(preview.pdf_dir, preview.workbook):
            raise ValueError("预览后输入文件或目录发生变化，请重新预览")
        if not preview.has_changes:
            return dict(status="no_change", total=preview.total, completed=[], cleanup="完成", payload=preview.payload)
        self.progress("正在检查 Excel 与文件写入权限…")
        self.excel.check(preview.workbook)
        pending_state = self._save_pending(preview)
        result = dict(status="failed", completed=[], total=preview.total, payload=preview.payload, engine_version=preview.engine_version, corrections=preview.corrections, original_payload=preview.original_payload, rows_filled=0)
        with self._scratch() as scratch:
            try:
                records = scratch / "preview.json"
                records.write_text(json.dumps(preview.payload, ensure_ascii=False), encoding="utf-8")
                writer_plan = plan_workbook(preview.workbook, preview.payload, template_rules=load_template_rules(self.engine_dir))
                skip_write = preview.resume_pending and pending_state.get("workbook_written") and writer_plan.workbook_matches
                if preview.fingerprints != self._fingerprints(preview.pdf_dir, preview.workbook):
                    writer_plan.workbook.close()
                    raise ValueError("执行前文件再次发生变化，请重新预览")
                self.progress("正在重命名发票…")
                rename_report = scratch / "rename-report.json"
                try:
                    renamed = self._invoke("pipeline", ["rename", "--records", str(records), "--apply", "--report", str(rename_report)], rename_report)
                    result["rename_plan"] = self._actual_rename_plan(preview)
                    if (renamed.get("exceptions") or len(renamed.get("rename_plan", [])) != len(preview.payload["rename_plan"])
                            or any(p.get("status") not in ("renamed", "unchanged") for p in renamed.get("rename_plan", []))
                            or any(p["status"] not in ("renamed", "unchanged") or sha256(p["destination"]) != preview.fingerprints[p["source"]] for p in result["rename_plan"])):
                        raise ValueError("发票改名未完成：" + json.dumps(renamed.get("exceptions", []), ensure_ascii=False))
                    result["completed"].append("发票改名")
                    if not skip_write:
                        self.progress("正在清空两表旧明细并按完整批次重填…")
                        write_workbook(writer_plan)
                        result["rows_filled"] = preview.rows_to_fill
                        result["detail_rows_removed"] = preview.blank_rows_to_remove
                        self._save_pending(preview, workbook_written=True)
                finally:
                    writer_plan.workbook.close()
                result["completed"].append("两表重填已在上次完成" if skip_write else "清空旧明细并重填两表")
                self.progress("正在应用字体、调整单元格完整显示并通过 Excel 重算…")
                self.excel.apply(self.engine_dir / "apply_mixed_fonts_excel.ps1", preview.workbook)
                result["completed"].append("字体、完整显示排版与重算")
                self.progress("正在逐票核对金额与两张表的总计…")
                totals_path = scratch / "totals.json"
                result["totals"] = self._invoke("totals", ["--input", str(preview.workbook), "--output", str(totals_path)], totals_path)
                verification_path = scratch / "verification.json"
                verified = self._invoke("pipeline", ["verify", "--records", str(records), "--excel-totals", str(totals_path), "--report", str(verification_path)], verification_path)
                result["verification"] = verified
                checks = verified.get("verification", [])
                if (verified.get("exceptions") or len(checks) != len(preview.payload["records"])
                        or any(v.get("status") != "passed" for v in checks)
                        or verified.get("batch_verification", {}).get("status") != "passed"):
                    raise ValueError("逐分金额验证失败，请查看异常和实际已执行步骤")
                result["completed"].append("逐分验证")
                self._pending_path(preview.pdf_dir, preview.workbook).unlink()
                result["status"] = "completed"
            except Exception as exc:
                result["error"] = str(exc)
                result["rename_plan"] = self._actual_rename_plan(preview)
        result["cleanup"] = self.cleanup_error or "完成"
        self.progress("处理完成" if result["status"] == "completed" else "处理未完成，请查看结果")
        return result
