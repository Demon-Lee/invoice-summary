"""Software-owned worksheet geometry shared by formatting and PDF export.

Only meaningful, nonempty cell anchors determine the range. Style-only cells
and retained blank detail rows must not move the final product row.
"""
from decimal import Decimal, InvalidOperation
import re

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import column_index_from_string, get_column_letter

from .workbook import HEADERS, SIGNERS, text


def sheet_layout(ws):
    """Return validated geometry without changing any worksheet property."""
    name = ws.title
    if name not in SIGNERS:
        raise ValueError(f"不支持的排版工作表：{name}")
    if ws.sheet_state != "visible":
        raise ValueError(f"{name} 工作表被隐藏，无法证明完整打印")
    cells = {(c.row, c.column): c for c in ws._cells.values()
             if not isinstance(c, MergedCell) and text(c.value)}
    if not cells:
        raise ValueError(f"{name} 没有可排版的内容")
    required = HEADERS + (("生产厂家",) if name == "验收" else ())
    rows = {}
    for cell in cells.values():
        value = text(cell.value)
        key = "生产厂家" if value.startswith("生产厂家") else value
        if key in required:
            rows.setdefault(cell.row, {}).setdefault(key, []).append(cell.column)
    headers = [(r, cols) for r, cols in rows.items() if all(k in cols for k in required)]
    if len(headers) != 1 or any(len(cols) != 1 for cols in headers[0][1].values()):
        raise ValueError(f"{name} 需要唯一完整且无重复的表头")
    header, mapping = headers[0]
    columns = {key: cols[0] for key, cols in mapping.items()}
    footers = [r for (r, c), cell in cells.items()
               if c == 1 and r > header and text(cell.value).startswith("合计")]
    if len(footers) != 1:
        raise ValueError(f"{name} 需要唯一的合计页脚")
    footer = footers[0]

    def value(row, column):
        cell = cells.get((row, column))
        return text(cell.value) if cell else ""

    first, second = SIGNERS[name]
    for label, column in ((first, 1), (second, 4)):
        roles = [c for (r, c), cell in cells.items()
                 if r == footer + 1 and text(cell.value).startswith(label)]
        if roles != [column]:
            raise ValueError(f"{name} 的签字行需要唯一的“{label}”标签")
    if not value(footer + 2, 1).startswith("说明"):
        raise ValueError(f"{name} 的“合计—签字—说明”三行页脚不完整")
    date_label = "验收日期" if name == "验收" else "领用日期"
    dates = [cell for cell in cells.values()
             if text(cell.value).rstrip("：:") == date_label]
    if (len(dates) != 1 or dates[0].row != footer + 1 or dates[0].column >= 16384
            or not value(footer + 1, dates[0].column + 1)):
        raise ValueError(f"{name} 需要唯一的日期标签及右侧已填写的日期单元格")
    detail_rows = []
    candidate_rows = sorted({r for r, c in cells if header < r < footer
                             and c in (columns["序号"], columns["产品名称"])})
    for row in candidate_rows:
        sequence, product = value(row, columns["序号"]), value(row, columns["产品名称"])
        try:
            number = Decimal(sequence)
            valid = number.is_finite() and number > 0 and number == number.to_integral_value()
        except (InvalidOperation, ValueError):
            valid = False
        if not valid or not product:
            raise ValueError(f"{name} 第 {row} 行需要有效序号及产品名称")
        detail_rows.append(row)
    if not detail_rows:
        raise ValueError(f"{name} 没有已填写的明细，请先完成填表")
    repeat_start, repeat_end = 1, header
    if ws.print_title_rows:
        match = re.fullmatch(r"\$?(\d+):\$?(\d+)", str(ws.print_title_rows))
        if not match:
            raise ValueError(f"{name} 重复打印标题行范围无效")
        repeat_start, repeat_end = map(int, match.groups())
        if not 1 <= repeat_start <= repeat_end <= header:
            raise ValueError(f"{name} 重复打印标题行必须位于表头行及其上方")
    last_row = max(r for r, _ in cells)
    last_col = max(c for _, c in cells)
    merges = list(ws.merged_cells.ranges)
    for index, merge in enumerate(merges):
        if not (1 <= merge.min_row <= merge.max_row <= 1048576
                and 1 <= merge.min_col <= merge.max_col <= 16384):
            raise ValueError(f"{name} 合并单元格范围无效：{merge}")
        for other in merges[:index]:
            if (merge.min_row <= other.max_row and other.min_row <= merge.max_row
                    and merge.min_col <= other.max_col and other.min_col <= merge.max_col):
                raise ValueError(f"{name} 存在重叠合并单元格：{other}、{merge}")
        if (merge.min_row, merge.min_col) in cells:
            last_row, last_col = max(last_row, merge.max_row), max(last_col, merge.max_col)
    if not (1 <= last_row <= 1048576 and 1 <= last_col <= 16384):
        raise ValueError(f"{name} 内容范围超出 Excel 可打印范围")
    if ws.sheet_format.zeroHeight:
        raise ValueError(f"{name} 设置了默认隐藏行，无法证明内容可见")
    for row, dimension in ws.row_dimensions.items():
        if 1 <= row <= last_row and (dimension.hidden or dimension.height == 0):
            raise ValueError(f"{name} 第 {row} 行被隐藏或高度为零")
    for key, dimension in ws.column_dimensions.items():
        start = dimension.min or column_index_from_string(key)
        end = dimension.max or start
        if start <= last_col and end >= 1 and (dimension.hidden or dimension.width == 0):
            raise ValueError(f"{name} {key} 列被隐藏或宽度为零")
    # Unmerged template captions intentionally spill into empty neighbours.
    # Preserve that existing display space without merging or moving cells.
    main_end = max(c for r, c in cells if r == header)
    caption_spans = {}
    for (row, column), cell in cells.items():
        if (header <= row < footer or column > main_end or cell.data_type not in ('s', 'inlineStr')
                or cell.alignment.horizontal not in (None, 'left', 'general')
                or any(cell.coordinate in merge for merge in merges)):
            continue
        end = min([main_end] + [c - 1 for r, c in cells if r == row and c > column])
        for merge in merges:
            if merge.min_row <= row <= merge.max_row and column < merge.min_col <= end:
                end = merge.min_col - 1
        if end > column:
            caption_spans[cell.coordinate] = f'{cell.coordinate}:{get_column_letter(end)}{row}'
    return {
        "name": name, "header_row": header, "footer_row": footer,
        "last_detail_row": detail_rows[-1], "print_end_row": last_row,
        "print_end_column": last_col, "print_area": f"A1:{get_column_letter(last_col)}{last_row}",
        "numeric_columns": [columns[k] for k in ("序号", "数量", "单价", "金额")],
        "general_cells": [cells[key].coordinate for key in sorted(cells) if cells[key].number_format == "General"],
        "content_cells": [cells[key].coordinate for key in sorted(cells)],
        "columns": columns, "tail_start_row": detail_rows[-1], "tail_end_row": last_row,
        "repeat_start_row": repeat_start, "repeat_end_row": repeat_end,
        "caption_spans": caption_spans,
    }


def layout_plan(path):
    """Load and close a saved workbook while producing the shared JSON plan."""
    wb = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    try:
        missing = set(SIGNERS) - set(wb.sheetnames)
        if missing:
            raise ValueError("工作簿缺少工作表：" + "、".join(sorted(missing)))
        return {"schema_version": 1, "sheets": [sheet_layout(wb[name]) for name in SIGNERS]}
    finally:
        wb.close()
