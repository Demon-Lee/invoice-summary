"""Versioned, offline snapshots of the user's invoice skill."""
import ast
import contextlib
import hashlib
import importlib
import importlib.util
import inspect
import io
import json
import os
import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from .policy import (SHARED_FILES, LEGACY_SHARED_FILES, shared_scope, validate_scope, snapshot_files,
                     LEGACY_WRITER_SHA256, TEMPLATE_WRITER_SHA256)
from .template_rules import read_template_rules

FILES = SHARED_FILES
MAX_BYTES = 8 * 1024 * 1024


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def source_files(skill_dir):
    root = Path(skill_dir).resolve()
    names = SHARED_FILES if (root / "scripts/template_rules.json").is_file() else LEGACY_SHARED_FILES
    result = {name: root / name if name == "SKILL.md" else root / "scripts" / name for name in names}
    writer = root / "scripts/write_summary.mjs"
    for path in (*result.values(), writer):
        if not path.is_file():
            raise ValueError(f"skill 文件不完整：{path}")
    return result, writer


def manifest_for_skill(skill_dir):
    files, writer = source_files(skill_dir)
    hashes = {name: sha256(path) for name, path in files.items()}
    marker = sha256(writer)
    format_version = 3 if "template_rules.json" in files else 2
    if format_version == 3:
        read_template_rules(files["template_rules.json"])
    manifest = {"format_version": format_version, "files": hashes, "source_writer_sha256": marker, "scope": shared_scope(format_version)}
    manifest["version"] = _snapshot_version(manifest)
    return manifest


def _snapshot_version(manifest):
    values = [manifest["files"], manifest["source_writer_sha256"]]
    if manifest["format_version"] >= 2:
        values.append(manifest["scope"])
    return hashlib.sha256(json.dumps(values, sort_keys=True).encode()).hexdigest()[:12]


def read_manifest(directory):
    root = Path(directory)
    manifest = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    if not isinstance(manifest, dict):
        raise ValueError("更新包清单必须为对象")
    validate_scope(manifest)
    if (not isinstance(manifest.get("files"), dict) or set(manifest["files"]) != set(snapshot_files(manifest))
            or not re.fullmatch(r"[a-f0-9]{12}", str(manifest.get("version", "")))
            or not re.fullmatch(r"[a-f0-9]{64}", str(manifest.get("source_writer_sha256", "")))):
        raise ValueError("更新包格式不兼容，需要升级主程序")
    for name, digest in manifest["files"].items():
        if (not re.fullmatch(r"[a-f0-9]{64}", str(digest))
                or not (root / name).is_file() or sha256(root / name) != digest):
            raise ValueError(f"更新文件完整性检查失败：{name}")
    expected = _snapshot_version(manifest)
    if expected != manifest["version"]:
        raise ValueError("更新版本完整性检查失败")
    if manifest["format_version"] == 3:
        read_template_rules(root / "template_rules.json")
    return manifest


def validate_python(directory):
    required = {"invoice_pipeline.py": {"main": 1, "is_verification_report_pdf": 1},
                "collect_excel_totals.py": {"main": 1, "collect_excel_totals": 1}}
    modules = {}
    for name in ("invoice_pipeline.py", "collect_excel_totals.py"):
        source = (Path(directory) / name).read_text(encoding="utf-8-sig")
        tree = ast.parse(source, filename=name)
        functions = {node.name: node for node in tree.body if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))}
        for function in required[name]:
            node = functions.get(function)
            if not isinstance(node, ast.FunctionDef):
                raise ValueError(f"识别接口不兼容：{name} 缺少同步函数 {function}")
            positional = node.args.posonlyargs + node.args.args
            if ((not positional and node.args.vararg is None) or len(positional)-len(node.args.defaults) > 1
                    or any(default is None for default in node.args.kw_defaults)):
                raise ValueError(f"识别接口不兼容：{name} 的 {function} 无法接收单个参数")
        for node in ast.walk(tree):
            imported = [a.name for a in node.names] if isinstance(node, ast.Import) else (
                [node.module] if isinstance(node, ast.ImportFrom) and node.module and not node.level else [])
            for module in imported:
                try:
                    importlib.import_module(module)
                except Exception as exc:
                    raise ValueError(f"更新依赖 {module} 无法加载，请升级主程序：{exc}") from exc
        try:
            from .controller import load_engine
            loaded = load_engine(directory, name)
            for function, count in required[name].items():
                callable_value = getattr(loaded, function)
                if inspect.iscoroutinefunction(callable_value):
                    raise TypeError("不支持异步入口")
                inspect.signature(callable_value).bind(*([None] * count))
            modules[name] = loaded
        except (Exception, SystemExit) as exc:
            raise ValueError(f"识别接口无法加载：{name}：{exc}") from exc
    _validate_output_contract(modules)


def _validate_output_contract(modules):
    """Exercise all shared CLI contracts only against generated empty inputs."""
    from openpyxl import Workbook
    with tempfile.TemporaryDirectory(prefix="invoice-contract-") as temp:
        root = Path(temp)
        def invoke(module, args, filename):
            target = root / filename
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                status = module.main([*args, "--output" if filename == "totals.json" else "--report", str(target)])
            if status not in (None, 0):
                raise ValueError("入口返回失败")
            result = json.loads(target.read_text(encoding="utf-8"))
            if not isinstance(result, dict):
                raise ValueError("结果不是对象")
            return result
        try:
            parser = modules["invoice_pipeline.py"]
            preview = invoke(parser, ["preview", "--pdf-dir", str(root)], "preview.json")
            if not all(isinstance(preview.get(key), list) for key in ("records", "rename_plan", "exceptions")):
                raise ValueError("预览字段不完整")
            records = root / "records.json"
            records.write_text('{"records": [], "exceptions": []}', encoding="utf-8")
            renamed = invoke(parser, ["rename", "--records", str(records), "--apply"], "rename.json")
            if not all(isinstance(renamed.get(key), list) for key in ("records", "rename_plan", "exceptions")):
                raise ValueError("改名字段不完整")
            wb = Workbook()
            wb.remove(wb.active)
            for title, first, second in (("验收", "经办人", "验收人"), ("领用", "保管人", "领用人")):
                ws = wb.create_sheet(title)
                ws.append(["序号", "金额"])
                ws.append(["合计", 0])
                ws.cell(3, 1, first + "：")
                ws.cell(3, 4, second + "：")
                ws.cell(4, 1, "说明：接口检查")
            book = root / "generated.xlsx"
            wb.save(book)
            wb.close()
            totals = invoke(modules["collect_excel_totals.py"], ["--input", str(book)], "totals.json")
            if totals != {"sequences": {}, "acceptance_footer_total": "0.00", "issue_footer_total": "0.00"}:
                raise ValueError("金额汇总字段不兼容")
            verified = invoke(parser, ["verify", "--records", str(records), "--excel-totals", str(root / "totals.json")], "verify.json")
            if (not isinstance(verified.get("verification"), list) or not isinstance(verified.get("exceptions"), list)
                    or verified.get("batch_verification", {}).get("status") != "passed"):
                raise ValueError("验证字段不完整")
        except (Exception, SystemExit) as exc:
            raise ValueError(f"识别接口输出不兼容：{exc}") from exc


def check_application_contract(candidate, supported):
    validate_scope(candidate)
    expected = TEMPLATE_WRITER_SHA256 if candidate["format_version"] == 3 else LEGACY_WRITER_SHA256
    if candidate["source_writer_sha256"] != expected:
        raise ValueError("skill 写表组件发生了软件尚未支持的变化，须适配软件后再同步。受支持的模板规则文件可直接同步；人工校正和清空重填策略会保留。")


def changed_files(candidate, baseline):
    changed = [name for name in FILES if candidate["files"].get(name) != baseline["files"].get(name)]
    if candidate["source_writer_sha256"] != baseline["source_writer_sha256"]:
        changed.append("write_summary.mjs")
    if (candidate["format_version"] != baseline["format_version"]
            or candidate.get("scope") != baseline.get("scope")):
        changed.append("manifest.json")
    return changed


def check_compatible(candidate, baseline):
    check_application_contract(candidate, baseline)
    changed = [name for name in FILES if candidate["files"].get(name) != baseline["files"].get(name)]
    if changed == ["SKILL.md"]:
        raise ValueError("只有规则说明变化，尚无对应脚本改进；请完成 skill 脚本适配后再同步")
    return changed_files(candidate, baseline)


def export_update(skill_dir, output, baseline=None):
    output = Path(output).resolve()
    paths, writer = source_files(skill_dir)
    if output in [p.resolve() for p in (*paths.values(), writer)] or output.suffix.lower() != ".zip":
        raise ValueError("更新包必须是独立的 .zip 文件，不能覆盖 skill 源文件")
    manifest = manifest_for_skill(skill_dir)
    check_application_contract(manifest, baseline)
    if baseline is not None:
        check_compatible(manifest, baseline)
    with tempfile.TemporaryDirectory(prefix="invoice-export-") as temp:
        root = Path(temp)
        for name, source in paths.items():
            shutil.copyfile(source, root / name)
        validate_python(root)
    output.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, source in paths.items():
            archive.write(source, name)
        archive.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
    return manifest


class EngineStore:
    def __init__(self, root, bundled):
        self.root = Path(root).resolve()
        self.bundled = Path(bundled).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        read_manifest(self.bundled)

    def _state(self):
        path = self.root / "active.json"
        return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {"active": "bundled", "previous": None}

    def _path(self, version):
        if version == "bundled":
            return self.bundled
        if not re.fullmatch(r"[a-f0-9]{12}", str(version)):
            raise ValueError("识别版本配置损坏")
        return self.root / "engines" / version

    def active(self):
        directory = self._path(self._state()["active"])
        check_application_contract(read_manifest(directory), read_manifest(self.bundled))
        return directory

    def version(self):
        return read_manifest(self.active())["version"]

    def source_status(self, source):
        candidate = manifest_for_skill(source)
        baseline = read_manifest(self.active())
        try:
            check_application_contract(candidate, read_manifest(self.bundled))
            changed = check_compatible(candidate, baseline)
            error = None
        except ValueError as exc:
            changed = changed_files(candidate, baseline)
            if candidate["source_writer_sha256"] != baseline["source_writer_sha256"] and "write_summary.mjs" not in changed:
                changed.append("write_summary.mjs")
            error = str(exc)
        return {"changed": changed, "error": error, "version": candidate["version"]}

    def _activate(self, state):
        path = self.root / "active.json"
        staged = self.root / "active.staged.json"
        staged.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
        os.replace(staged, path)

    def install(self, archive_path):
        with tempfile.TemporaryDirectory(prefix="update-", dir=self.root) as temp:
            extracted = Path(temp) / "engine"
            extracted.mkdir()
            with zipfile.ZipFile(archive_path) as archive:
                names = archive.namelist()
                allowed = (set(LEGACY_SHARED_FILES) | {"manifest.json"}, set(SHARED_FILES) | {"manifest.json"})
                if set(names) not in allowed or len(names) != len(set(names)):
                    raise ValueError("更新包包含不允许的路径或缺少必要文件")
                if sum(item.file_size for item in archive.infolist()) > MAX_BYTES:
                    raise ValueError("更新包超过允许大小")
                for item in archive.infolist():
                    if item.is_dir() or (item.external_attr >> 16) & 0o170000 == 0o120000:
                        raise ValueError("更新包包含不允许的文件类型")
                    (extracted / item.filename).write_bytes(archive.read(item))
            manifest = read_manifest(extracted)
            if set(names) != set(snapshot_files(manifest)) | {"manifest.json"}:
                raise ValueError("更新包文件与其声明的规则格式不一致")
            before = read_manifest(self.active())
            check_application_contract(manifest, read_manifest(self.bundled))
            changed = check_compatible(manifest, before)
            validate_python(extracted)
            destination = self.root / "engines" / manifest["version"]
            destination.parent.mkdir(exist_ok=True)
            if not destination.exists():
                os.replace(extracted, destination)
            else:
                read_manifest(destination)
            state = self._state()
            if state["active"] != manifest["version"]:
                self._activate({"active": manifest["version"], "previous": state["active"]})
            return {"version": manifest["version"], "previous_version": before["version"], "changed": changed}

    def sync(self, source):
        with tempfile.TemporaryDirectory(prefix="sync-", dir=self.root) as temp:
            path = Path(temp) / "update.zip"
            export_update(source, path, read_manifest(self.active()))
            return self.install(path)

    def rollback(self):
        state = self._state()
        if state.get("previous") is None:
            raise ValueError("没有可回退的识别版本")
        before = read_manifest(self.active())
        candidate = read_manifest(self._path(state["previous"]))
        check_application_contract(candidate, read_manifest(self.bundled))
        self._activate({"active": state["previous"], "previous": state["active"]})
        return {"version": candidate["version"], "previous_version": before["version"], "changed": changed_files(candidate, before)}

